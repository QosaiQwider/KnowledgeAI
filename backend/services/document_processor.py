from pypdf import PdfReader
from docx import Document as DocxDocument
from pptx import Presentation


# =========================================================
# EXTRACT TEXT
# =========================================================

def extract_text(file_path: str, file_type: str):

    file_type = file_type.lower()

    if file_type == "pdf":
        return extract_pdf(file_path)

    elif file_type == "docx":
        return extract_docx(file_path)

    elif file_type == "pptx":
        return extract_pptx(file_path)

    elif file_type == "txt":
        return extract_txt(file_path)

    else:
        raise ValueError("Unsupported file type")


# =========================================================
# PDF
# =========================================================

def extract_pdf(file_path):

    reader = PdfReader(file_path)

    pages = []

    for page_number, page in enumerate(reader.pages, start=1):

        text = page.extract_text() or ""

        if text.strip():
            pages.append({
                "page_number": page_number,
                "text": text
            })

    return pages


# =========================================================
# DOCX
# =========================================================

def extract_docx(file_path):

    document = DocxDocument(file_path)

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    )

    return [{
        "page_number": None,
        "text": text
    }]


# =========================================================
# PPTX
# =========================================================

def extract_pptx(file_path):

    presentation = Presentation(file_path)

    pages = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        texts = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):
                if shape.text.strip():
                    texts.append(shape.text)

        text = "\n".join(texts)

        if text.strip():
            pages.append({
                "page_number": slide_number,
                "text": text
            })

    return pages


# =========================================================
# TXT
# =========================================================

def extract_txt(file_path):

    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore"
    ) as file:

        text = file.read()

    return [{
        "page_number": None,
        "text": text
    }]


# =========================================================
# CHUNK TEXT
# =========================================================

def chunk_text(
    text: str,
    chunk_size: int = 1000,
    overlap: int = 200
):

    chunks = []

    start = 0

    while start < len(text):

        end = start + chunk_size

        chunk = text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        if end >= len(text):
            break

        start = end - overlap

    return chunks